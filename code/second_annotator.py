"""Second-annotator workflow for the two human audits.

  python code/second_annotator.py make    # writes annotation2/second_annotator_sheets.xlsx (blind) and annotation2/key.csv
  python code/second_annotator.py score   # reads the filled workbook, writes ann2 into both sheets, prints agreement

Blinding: items are shuffled; the observer, the lexical score, and the first annotator's labels are hidden.
"""
import csv, json, random, sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
IND = ROOT / "results/induction/annotation_sheet.csv"
FID = ROOT / "results/mr_fidelity/fidelity_sheet.csv"
OUT = ROOT / "annotation2"
WB = OUT / "second_annotator_sheets.xlsx"
KEY = OUT / "key.csv"

def read(p):
    with open(p, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))

def write(p, rows):
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)

INSTR_IND = ("For each item, read the benchmark secret and the observer's candidate list. Label M if at least one "
             "candidate states the same fact as the secret (any wording or language), A if a candidate is related but "
             "incomplete (e.g., right person, only part of the fact), X otherwise. Compare the candidates with the "
             "benchmark secret only, and do not consult other labels, AI tools, or the other sheet.")
INSTR_FID = ("For each item, compare the original and transformed excerpt. Label Y if the transformation is faithful to "
             "its question (MR1: only wording changed; MR3: inserted sentences unrelated to the dialogue and secret; "
             "MR4: only person names changed, consistently, including in the secret), N otherwise. Do not use AI tools.")

def make():
    OUT.mkdir(exist_ok=True)
    rng = random.Random(20261006)
    ind = [r for r in read(IND) if r.get("final")]
    fid = read(FID)
    ind_items = [("I", i, r) for i, r in enumerate(ind)]
    fid_items = [("F", i, r) for i, r in enumerate(fid)]
    rng.shuffle(ind_items); rng.shuffle(fid_items)
    wb = Workbook()
    ws = wb.active; ws.title = "induction"
    ws.append([INSTR_IND]); ws.append(["item", "benchmark_secret", "candidates", "label (M/A/X)"])
    key = []
    for n, (_, i, r) in enumerate(ind_items, 1):
        iid = f"I{n:03d}"
        ws.append([iid, r["benchmark_secret"], r["candidates"].replace(" || ", "\n"), ""])
        key.append({"item": iid, "sheet": "induction", "sid": r["sid"], "observer": r["observer"], "mr": ""})
    dv = DataValidation(type="list", formula1='"M,A,X"', allow_blank=True); ws.add_data_validation(dv)
    dv.add(f"D3:D{len(ind_items)+2}")
    ws2 = wb.create_sheet("transformations")
    ws2.append([INSTR_FID]); ws2.append(["item", "mr", "question", "original", "transformed", "secret", "label (Y/N)"])
    for n, (_, i, r) in enumerate(fid_items, 1):
        iid = f"F{n:03d}"
        ws2.append([iid, r["mr"], r["question"], r["original"], r["transformed"], r["secret"], ""])
        key.append({"item": iid, "sheet": "transformations", "sid": r["sid"], "observer": "", "mr": r["mr"]})
    dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True); ws2.add_data_validation(dv2)
    dv2.add(f"G3:G{len(fid_items)+2}")
    for s, widths in ((ws, [8, 50, 80, 14]), (ws2, [8, 6, 30, 60, 60, 40, 12])):
        for c, w in zip("ABCDEFG", widths): s.column_dimensions[c].width = w
        for row in s.iter_rows(min_row=3):
            for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
        s["A1"].font = Font(bold=True)
    wb.save(WB); write(KEY, key)
    print(f"wrote {WB} ({len(ind_items)} induction, {len(fid_items)} transformation items) and {KEY}")

def kappa(a, b, cats):
    n = len(a); po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return po, (po - pe) / (1 - pe) if pe < 1 else float("nan")

def score():
    wb = load_workbook(WB)
    key = {k["item"]: k for k in read(KEY)}
    lab = {}
    for s in ("induction", "transformations"):
        for row in wb[s].iter_rows(min_row=3, values_only=True):
            if row[0]: lab[row[0]] = (str(row[-1] or "").strip().upper())
    missing = [i for i, v in lab.items() if not v]
    if missing: sys.exit(f"{len(missing)} items unlabeled, e.g. {missing[:5]}")
    ind = read(IND); fid = read(FID)
    idx_i = {(r["sid"], r["observer"]): r for r in ind}
    idx_f = {(r["sid"], r["mr"]): r for r in fid}
    for iid, v in lab.items():
        k = key[iid]
        if k["sheet"] == "induction": idx_i[(k["sid"], k["observer"])]["ann2"] = v
        else: idx_f[(k["sid"], k["mr"])]["ann2"] = v
    write(IND, ind); write(FID, fid)
    a1 = [r["final"] for r in ind if r.get("final")]; a2 = [r["ann2"] for r in ind if r.get("final")]
    po3, k3 = kappa(a1, a2, ["M", "A", "X"])
    b1 = ["M" if x == "M" else "N" for x in a1]; b2 = ["M" if x == "M" else "N" for x in a2]
    po2, k2 = kappa(b1, b2, ["M", "N"])
    f1 = [r["final"] for r in fid]; f2 = [r["ann2"] for r in fid]
    pof, kf = kappa(f1, f2, ["Y", "N"])
    res = {"induction_n": len(a1), "induction_agreement_3class": po3, "induction_kappa_3class": k3,
           "induction_agreement_M_vs_not": po2, "induction_kappa_M_vs_not": k2,
           "fidelity_n": len(f1), "fidelity_agreement": pof, "fidelity_kappa": kf,
           "fidelity_ann2_yes": f2.count("Y")}
    (OUT / "agreement.json").write_text(json.dumps(res, indent=2))
    print(json.dumps(res, indent=2))

if __name__ == "__main__":
    {"make": make, "score": score}[sys.argv[1] if len(sys.argv) > 1 else "make"]()
